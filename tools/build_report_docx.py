from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\ADMIN\Downloads\Template bài tập lớn Trí tuệ nhân tạo (1).docx")
OUTPUT = ROOT / "outputs" / "Bao_cao_nhap_DrowsyDriverAndroid.docx"


def main() -> None:
    doc = Document(TEMPLATE)
    clear_body(doc)
    configure_styles(doc)

    add_title(doc)
    add_assignment_table(doc)
    add_toc_placeholder(doc)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_conclusion(doc)
    add_references(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("HỌC VIỆN NGÂN HÀNG\nKHOA HỆ THỐNG THÔNG TIN QUẢN LÝ")
    run.bold = True
    run.font.size = Pt(13)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BÀI TẬP LỚN CUỐI KỲ\nMÔN: TRÍ TUỆ NHÂN TẠO")
    run.bold = True
    run.font.size = Pt(16)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Hệ thống phát hiện và cảnh báo dấu hiệu buồn ngủ của tài xế\n"
        "trên thiết bị Android sử dụng MediaPipe và mạng nơ-ron tích chập"
    )
    run.bold = True
    run.font.size = Pt(15)

    doc.add_paragraph()
    for line in [
        "Nhóm sinh viên: ................................................",
        "Lớp học phần: IS54A",
        "Giảng viên hướng dẫn: ..........................................",
        "Năm học: 2025 - 2026",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_assignment_table(doc: Document) -> None:
    doc.add_heading("Bảng phân công công việc", level=1)
    table = doc.add_table(rows=1, cols=4)
    apply_table_style(table)
    headers = ["Họ tên - Mã sinh viên", "Quyển báo cáo", "Sản phẩm", "Tỷ lệ đóng góp"]
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text

    rows = [
        [
            "Thành viên 1",
            "Chương 1; Chương 3",
            "Android app, MediaPipe pipeline, EAR/MAR, cảnh báo",
            "TBD",
        ],
        [
            "Thành viên 2",
            "Chương 2; Chương 4",
            "Dataset, train CNN, đánh giá, confusion matrix",
            "TBD",
        ],
        [
            "Thành viên 3",
            "Kết luận; tài liệu tham khảo",
            "Demo, video backup, Q&A thuyết trình",
            "TBD",
        ],
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text


def add_toc_placeholder(doc: Document) -> None:
    doc.add_heading("Mục lục", level=1)
    doc.add_paragraph("Cập nhật mục lục tự động trong Word sau khi hoàn thiện báo cáo.")


def add_chapter_1(doc: Document) -> None:
    doc.add_heading("Chương 1: Giới thiệu bài toán", level=1)
    doc.add_heading("1.1. Phát biểu bài toán", level=2)
    add_paragraphs(
        doc,
        [
            "Lái xe trong trạng thái buồn ngủ là một trong những nguyên nhân gây tai nạn giao thông nghiêm trọng. Các dấu hiệu như mắt nhắm lâu, ngáp liên tục, cúi đầu hoặc mất tập trung có thể được quan sát thông qua camera đặt trong cabin.",
            "Đề tài xây dựng prototype phát hiện và cảnh báo dấu hiệu buồn ngủ của tài xế bằng thiết bị Android phone. Input của hệ thống là luồng hình ảnh realtime từ camera Android. Output là trạng thái tài xế gồm No face, Awake, Eyes closed, Yawning và Drowsy alert.",
            "Phạm vi của đề tài là một tài xế trong khung hình, camera đặt gần chính diện khuôn mặt, xử lý trên thiết bị Android. Hệ thống là prototype AI có tính ứng dụng và có lộ trình phát triển thành sản phẩm thực tế.",
        ],
    )
    doc.add_heading("1.2. Ứng dụng của bài toán", level=2)
    add_bullets(
        doc,
        [
            "Hỗ trợ cá nhân lái xe đường dài, tài xế taxi, xe công nghệ, xe giao hàng, xe khách hoặc xe tải.",
            "Tận dụng Android phone có sẵn camera, loa, rung, CPU/GPU và màn hình.",
            "Xử lý on-device, không phụ thuộc internet, giảm độ trễ và phù hợp hơn với bối cảnh xe đang di chuyển.",
            "Có giá trị cho doanh nghiệp vận tải, nhóm nghiên cứu edge AI và cộng đồng an toàn giao thông.",
        ],
    )
    doc.add_heading("1.3. Khảo sát các bài làm liên quan", level=2)
    add_paragraphs(
        doc,
        [
            "Các hướng tiếp cận phổ biến gồm: hệ thống rule-based dùng EAR/MAR/PERCLOS/head pose; CNN hoặc deep learning phân loại trạng thái mắt, miệng hoặc khuôn mặt; và hệ thống hybrid kết hợp landmark với CNN/TFLite.",
            "Khoảng trống nghiên cứu chính là nhiều demo chỉ đánh giá offline, ít quan tâm latency/FPS trên thiết bị thật; nhiều hệ thống laptop webcam khó triển khai vào xe; và nhiều bài chỉ dùng accuracy trong khi bài toán an toàn cần chú trọng recall của lớp buồn ngủ.",
            "Đề tài này tập trung vào Android phone như thiết bị edge chi phí thấp, kết hợp landmark có thể giải thích với CNN/TFLite, đồng thời dùng smoothing theo thời gian để giảm báo nhầm do chớp mắt tự nhiên.",
        ],
    )


def add_chapter_2(doc: Document) -> None:
    doc.add_heading("Chương 2: Chuẩn bị dữ liệu", level=1)
    doc.add_heading("2.1. Thu thập dữ liệu", level=2)
    add_paragraphs(
        doc,
        [
            "Dữ liệu dự kiến gồm ba nhóm: dataset public liên quan đến tài xế buồn ngủ như NTHU Driver Drowsiness Detection Dataset, YawDD và DROZY; dataset eye-state/yawning để train nhanh; và dữ liệu tự thu bằng Android phone để kiểm tra điều kiện camera thực tế.",
            "Bảng thống kê số lượng mẫu theo lớp sẽ được điền sau khi hoàn thành tải dữ liệu, trích frame và crop ROI.",
        ],
    )
    add_simple_table(
        doc,
        ["Nguồn dữ liệu", "Số mẫu", "Nhãn", "Vai trò"],
        [
            ["Dataset public", "TBD", "eyes_open/eyes_closed/yawning", "Train/validation"],
            ["Video tự quay bằng Android", "TBD", "awake/closed/yawn", "Test/demo"],
        ],
    )
    doc.add_heading("2.2. Tiền xử lý dữ liệu", level=2)
    add_bullets(
        doc,
        [
            "Phát hiện khuôn mặt và landmark bằng MediaPipe.",
            "Crop vùng mắt hoặc miệng theo landmark.",
            "Resize ROI về 64x64 hoặc 96x96.",
            "Normalize pixel về [0, 1].",
            "Chia train/validation/test; nếu có thể chia theo subject/video để tránh rò rỉ dữ liệu.",
            "Dùng augmentation nhẹ như thay đổi độ sáng, xoay nhỏ và zoom nhỏ.",
        ],
    )


def add_chapter_3(doc: Document) -> None:
    doc.add_heading("Chương 3: Xây dựng mô hình", level=1)
    doc.add_heading("3.1. Trích chọn đặc trưng", level=2)
    add_paragraphs(
        doc,
        [
            "Hệ thống sử dụng đặc trưng hình học từ landmark gồm EAR cho mắt, MAR cho miệng, thời gian mắt nhắm liên tục và thời gian ngáp liên tục. Ngoài ra, hệ thống crop eye ROI để CNN học đặc trưng ảnh của vùng mắt.",
        ],
    )
    doc.add_heading("3.2. Lựa chọn thuật toán/mô hình", level=2)
    add_paragraphs(
        doc,
        [
            "Pipeline cuối gồm CameraX Android, MediaPipe Face Landmarker LIVE_STREAM, tính EAR/MAR, crop vùng mắt, CNN/TFLite phân loại mắt mở/mắt nhắm, temporal smoothing và cảnh báo âm thanh/rung.",
            "Baseline là MediaPipe + EAR/MAR threshold. Thành phần học sâu là CNN nhẹ, export sang TensorFlow Lite để chạy trên Android.",
        ],
    )
    doc.add_heading("3.3. Cấu hình huấn luyện mô hình", level=2)
    add_simple_table(
        doc,
        ["Thông số", "Giá trị đề xuất"],
        [
            ["Input CNN", "64x64x3 eye ROI"],
            ["Lớp", "eyes_open, eyes_closed"],
            ["Optimizer", "Adam"],
            ["Loss", "Categorical crossentropy"],
            ["Epoch", "10-20"],
            ["Batch size", "32"],
            ["Deployment", "TensorFlow Lite trên Android"],
        ],
    )


def add_chapter_4(doc: Document) -> None:
    doc.add_heading("Chương 4: Đánh giá mô hình và demo", level=1)
    doc.add_heading("4.1. Độ đo đánh giá", level=2)
    add_paragraphs(
        doc,
        [
            "Độ đo đánh giá gồm accuracy, precision, recall, F1-score, confusion matrix và FPS/latency khi chạy demo. Trong bài toán an toàn, recall của lớp buồn ngủ hoặc mắt nhắm quan trọng hơn accuracy tổng, vì bỏ sót tài xế buồn ngủ nguy hiểm hơn báo nhầm.",
        ],
    )
    doc.add_heading("4.2. Kết quả đánh giá", level=2)
    add_paragraphs(
        doc,
        [
            "Kết quả được tạo bằng script tools/evaluate_eye_classifier.py. Script đọc tập test, chạy model Keras hoặc TensorFlow Lite, sau đó xuất metrics.json, confusion_matrix.csv và predictions.csv.",
        ],
    )
    add_simple_table(
        doc,
        ["Phương pháp", "Accuracy", "Recall closed/drowsy", "F1-score", "FPS", "Nhận xét"],
        [
            ["EAR/MAR threshold", "TBD", "TBD", "TBD", "TBD", "Nhanh, dễ giải thích"],
            ["CNN eye classifier", "TBD", "TBD", "TBD", "TBD", "Học đặc trưng ảnh"],
            ["Hybrid + smoothing", "TBD", "TBD", "TBD", "TBD", "Ổn định hơn khi demo"],
        ],
    )
    doc.add_heading("4.3. Demo", level=2)
    add_bullets(
        doc,
        [
            "Ứng dụng Android mở camera phone.",
            "Overlay hiển thị trạng thái, EAR, MAR, FPS, confidence và nhãn CNN nếu model TFLite có mặt.",
            "Khi mắt nhắm liên tục hơn 1.7 giây, hệ thống chuyển sang Drowsy alert.",
            "Cảnh báo bằng âm thanh và rung.",
            "Có video demo backup để đảm bảo thuyết trình thuận lợi.",
        ],
    )


def add_conclusion(doc: Document) -> None:
    doc.add_heading("Kết luận", level=1)
    add_paragraphs(
        doc,
        [
            "Đề tài đã xây dựng pipeline prototype cho bài toán phát hiện dấu hiệu buồn ngủ của tài xế trên Android. Hướng tiếp cận kết hợp MediaPipe landmark, baseline EAR/MAR và CNN/TFLite giúp hệ thống vừa có khả năng giải thích, vừa có lộ trình triển khai AI trên thiết bị biên.",
            "Hạn chế hiện tại gồm phụ thuộc ánh sáng, góc camera, kính/che mặt; dataset tự thu còn nhỏ; cần kiểm thử trên nhiều người và nhiều thiết bị Android. Hướng phát triển là tích hợp đầy đủ CNN/TFLite, thêm PERCLOS, head pose, gaze, cá nhân hóa ngưỡng theo từng tài xế và lưu log sự kiện.",
        ],
    )
    doc.add_heading("Khía cạnh riêng tư và an toàn", level=2)
    add_paragraphs(
        doc,
        [
            "Vì hệ thống sử dụng camera hướng vào khuôn mặt tài xế, nhóm ưu tiên xử lý trực tiếp trên thiết bị và không lưu video khuôn mặt mặc định. Nếu triển khai trong tổ chức, hệ thống cần có thông báo rõ ràng, sự đồng ý của người dùng, chính sách lưu trữ dữ liệu và chỉ nên lưu log sự kiện tối thiểu thay vì video liên tục.",
            "Hệ thống chỉ là công cụ hỗ trợ cảnh báo sớm, không thay thế trách nhiệm nghỉ ngơi của tài xế và không phải hệ thống an toàn đã được chứng nhận cho xe thương mại. Trước khi dùng thực tế cần kiểm thử có kiểm soát trong nhiều điều kiện ánh sáng, góc camera và nhóm người dùng khác nhau.",
        ],
    )


def add_references(doc: Document) -> None:
    doc.add_heading("Tài liệu tham khảo", level=1)
    refs = [
        "Google AI Edge. MediaPipe Face Landmarker for Android. https://ai.google.dev/edge/mediapipe/solutions/vision/face_landmarker/android",
        "Android Developers. CameraX ImageAnalysis. https://developer.android.com/media/camera/camerax/analyze",
        "TensorFlow. TensorFlow Lite Delegates. https://www.tensorflow.org/lite/performance/delegates",
        "NTHU CVLab. Driver Drowsiness Detection Dataset. https://cv.cs.nthu.edu.tw/php/callforpaper/datasets/DDD/",
        "DROZY: The ULg Multimodality Drowsiness Database. https://www.drozy.uliege.be/",
        "A Real-Time Embedded System for Driver Drowsiness Detection Based on Visual Analysis of the Eyes and Mouth Using CNN and MAR. https://pmc.ncbi.nlm.nih.gov/articles/PMC11479241/",
    ]
    for index, ref in enumerate(refs, start=1):
        doc.add_paragraph(f"{index}. {ref}")


def add_paragraphs(doc: Document, paragraphs: list[str]) -> None:
    for text in paragraphs:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(f"- {item}")


def add_simple_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    apply_table_style(table)
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    doc.add_paragraph()


def apply_table_style(table) -> None:
    try:
        table.style = "Table Grid"
    except KeyError:
        pass


if __name__ == "__main__":
    main()
