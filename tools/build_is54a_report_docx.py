from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "outputs" / "Bao_cao_IS54A_DrowsyDriverAndroid_theo_khung.docx"


def main() -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    configure_styles(doc)

    add_cover(doc)
    add_assignment_table(doc)
    add_toc_and_abbreviations(doc)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_conclusion(doc)
    add_references(doc)

    doc.save(OUT_PATH)
    print(f"Saved {OUT_PATH}")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(31, 78, 121), 18, 8),
        ("Heading 2", 14, RGBColor(31, 78, 121), 14, 6),
        ("Heading 3", 12, RGBColor(67, 67, 67), 10, 4),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HỌC VIỆN: [ĐIỀN TÊN HỌC VIỆN THEO QUY ĐỊNH]\n")
    r.bold = True
    r.font.size = Pt(14)
    r = p.add_run("KHOA/BỘ MÔN: [ĐIỀN KHOA/BỘ MÔN]\n")
    r.bold = True
    r.font.size = Pt(13)

    add_spacer(doc, 3)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÀI TẬP LỚN CUỐI KỲ\n")
    r.bold = True
    r.font.size = Pt(18)
    r = p.add_run("Môn học: Trí tuệ nhân tạo (IS54A)")
    r.font.size = Pt(14)

    add_spacer(doc, 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TÊN ĐỀ TÀI\n")
    r.bold = True
    r.font.size = Pt(14)
    r = p.add_run(
        "Phát hiện và cảnh báo trạng thái buồn ngủ của tài xế "
        "trong phạm vi thiết bị Android/edge camera "
        "sử dụng MediaPipe Face Landmarker và mạng nơ-ron tích chập"
    )
    r.bold = True
    r.font.size = Pt(16)

    add_spacer(doc, 2)
    add_key_value(doc, "Tên bài toán", "Phát hiện và cảnh báo trạng thái buồn ngủ của tài xế.")
    add_key_value(doc, "Phạm vi", "Prototype chạy trên Android phone; định hướng triển khai thực tế trên camera hành trình thông minh hoặc thiết bị embedded trong xe.")
    add_key_value(doc, "Kỹ thuật sử dụng", "CameraX, MediaPipe Face Landmarker, EAR/MAR, CNN, TensorFlow/Keras, TensorFlow Lite và Streamlit fallback.")

    add_spacer(doc, 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Thành viên nhóm")
    r.bold = True
    r.font.size = Pt(13)

    members = [
        "[Họ tên 1] - [Mã sinh viên]",
        "[Họ tên 2] - [Mã sinh viên]",
        "[Họ tên 3] - [Mã sinh viên]",
        "[Họ tên 4] - [Mã sinh viên]",
    ]
    for member in members:
        p = doc.add_paragraph(member)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_spacer(doc, 3)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Hà Nội, [tháng] năm 2026").italic = True
    doc.add_page_break()


def add_assignment_table(doc: Document) -> None:
    doc.add_heading("Bảng phân công công việc, tỷ lệ phần trăm đóng góp", level=1)
    add_note(
        doc,
        "Phần trăm đóng góp tổng của các thành viên trong nhóm là 100%. "
        "Con số này mang tính tham khảo để giảng viên cân nhắc khi có điểm cộng/trừ, "
        "không mang tính tuyệt đối để tính điểm cuối cùng. "
        "Tất cả các section trong quyển báo cáo, trừ Mục lục và Tài liệu tham khảo, "
        "đều được khai báo trong bảng phân công dưới đây.",
    )

    rows = [
        [
            "[Họ tên 1] - [Mã sinh viên]",
            "Chương 1: 1.1. Phát biểu bài toán; 1.2. Ứng dụng của bài toán; 1.3. Khảo sát bài làm liên quan.\nKết luận: tổng kết và hạn chế.",
            "Tài liệu khảo sát; chỉnh nội dung báo cáo; chuẩn bị câu hỏi bảo vệ.\nFiles: docs/DE_CUONG_BAO_CAO.md, research/SOURCES_FOR_REPORT.md",
            "25%\n\n(Ký tên)",
        ],
        [
            "[Họ tên 2] - [Mã sinh viên]",
            "Chương 2: 2.1. Thu thập dữ liệu; 2.2. Tiền xử lý dữ liệu.",
            "Chuẩn hóa dataset; thống kê dữ liệu; chuẩn bị manifest.\nFiles: tools/prepare_eye_dataset.py, tools/summarize_dataset.py, outputs/dataset_summary/*",
            "25%\n\n(Ký tên)",
        ],
        [
            "[Họ tên 3] - [Mã sinh viên]",
            "Chương 3: 3.1. Trích chọn đặc trưng; 3.2. Lựa chọn thuật toán/mô hình; 3.3. Cấu hình huấn luyện mô hình.",
            "Train CNN; export TFLite; mô tả pipeline AI.\nFiles: tools/train_eye_classifier.py, app/src/main/assets/drowsiness_model.tflite",
            "25%\n\n(Ký tên)",
        ],
        [
            "[Họ tên 4] - [Mã sinh viên]",
            "Chương 4: 4.1. Độ đo đánh giá; 4.2. Kết quả đánh giá; 4.3. Demo.",
            "Android app, Streamlit fallback, evaluate model, video demo.\nFiles: app/src/main/java/com/ai2026/drowsydriver/*, tools/evaluate_eye_classifier.py, streamlit_app.py",
            "25%\n\n(Ký tên)",
        ],
    ]

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    set_table_widths(table, [2400, 3000, 3000, 1560])
    headers = ["Họ tên - Mã sinh viên", "Quyển báo cáo", "Sản phẩm", "Phần trăm đóng góp - Ký tên xác nhận"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade_cell(cell, "D9EAF7")
        set_cell_text_bold(cell)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    doc.add_page_break()


def add_toc_and_abbreviations(doc: Document) -> None:
    doc.add_heading("Mục lục", level=1)
    p = doc.add_paragraph()
    add_toc_field(p)
    add_note(doc, "Sau khi mở file trong Microsoft Word: bấm chuột phải vào mục lục và chọn Update Field để cập nhật số trang.")

    doc.add_page_break()
    doc.add_heading("Danh mục từ viết tắt", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_widths(table, [1600, 3000, 4960])
    for i, text in enumerate(["Từ viết tắt", "Tên đầy đủ", "Ý nghĩa trong đề tài"]):
        table.rows[0].cells[i].text = text
        shade_cell(table.rows[0].cells[i], "D9EAF7")
        set_cell_text_bold(table.rows[0].cells[i])
    abbreviations = [
        ("AI", "Artificial Intelligence", "Trí tuệ nhân tạo."),
        ("CNN", "Convolutional Neural Network", "Mạng nơ-ron tích chập dùng phân loại trạng thái mắt."),
        ("EAR", "Eye Aspect Ratio", "Tỷ lệ hình học dùng phát hiện mắt nhắm dựa trên landmarks."),
        ("MAR", "Mouth Aspect Ratio", "Tỷ lệ hình học dùng phát hiện há miệng/ngáp."),
        ("TFLite", "TensorFlow Lite", "Định dạng mô hình tối ưu cho thiết bị di động/edge."),
        ("ROI", "Region of Interest", "Vùng ảnh quan tâm, trong đề tài là vùng mắt."),
    ]
    for row in abbreviations:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    doc.add_page_break()


def add_chapter_1(doc: Document) -> None:
    doc.add_heading("Chương 1: Giới thiệu bài toán", level=1)
    doc.add_heading("1.1. Phát biểu", level=2)
    add_paragraphs(
        doc,
        [
            "Tên bài toán của đề tài là phát hiện và cảnh báo trạng thái buồn ngủ của tài xế. Hệ thống nhận đầu vào là khung hình từ camera hướng về khuôn mặt tài xế, sau đó phân tích trạng thái mắt và miệng để đưa ra cảnh báo khi có dấu hiệu buồn ngủ hoặc ngáp.",
            "Input của hệ thống gồm ảnh/video frame từ camera, landmarks khuôn mặt do MediaPipe Face Landmarker trích xuất và vùng ảnh mắt được crop từ frame. Output gồm trạng thái Awake, Drowsy hoặc Yawning; đồng thời app hiển thị overlay, phát cảnh báo âm thanh/rung và ghi log sự kiện tối thiểu.",
            "Dạng dữ liệu chính là ảnh khuôn mặt/ảnh mắt. Trong bản prototype, Android phone được dùng như thiết bị edge thay thế cho camera hành trình thông minh, vì điện thoại có sẵn camera, bộ xử lý, màn hình, loa và khả năng chạy TensorFlow Lite.",
        ],
    )
    add_placeholder(doc, "Chèn ảnh minh họa input: frame khuôn mặt tài xế hoặc ảnh mắt.")
    add_placeholder(doc, "Chèn ảnh minh họa output: giao diện app hiển thị Awake/Drowsy/Yawning.")

    doc.add_heading("1.2. Ứng dụng của bài toán", level=2)
    add_paragraphs(
        doc,
        [
            "Bài toán phát hiện buồn ngủ có ý nghĩa trong an toàn giao thông, đặc biệt với tài xế lái xe đường dài, xe khách, xe tải hoặc taxi công nghệ. Hệ thống có thể cảnh báo sớm khi tài xế nhắm mắt kéo dài, ngáp liên tục hoặc mất tập trung.",
            "Trong thực tế, hệ thống có thể tích hợp vào camera hành trình thông minh, thiết bị ADAS gắn trên xe hoặc hộp xử lý embedded. Trong phạm vi bài tập lớn, đề tài giới hạn ở prototype chạy trên điện thoại Android và một bản Streamlit fallback để minh họa thuật toán trên máy tính.",
            "Phạm vi người dùng mục tiêu là tài xế cá nhân hoặc tài xế dịch vụ cần cảnh báo cơ bản. Hệ thống chưa được chứng nhận an toàn giao thông, chưa thay thế thiết bị giám sát chuyên dụng và chưa đảm bảo hoạt động tốt trong mọi điều kiện ánh sáng, kính râm hoặc góc mặt khó.",
        ],
    )

    doc.add_heading("1.3. Khảo sát các bài làm liên quan", level=2)
    add_paragraphs(
        doc,
        [
            "Các hướng nghiên cứu phổ biến gồm dùng landmark khuôn mặt để tính EAR/MAR, dùng CNN phân loại mắt mở/mắt nhắm, dùng YOLO để phát hiện trực tiếp mắt mở/mắt nhắm/ngáp, hoặc dùng dữ liệu đa phương thức như video, tín hiệu sinh lý và thang đo buồn ngủ.",
            "So với hướng YOLO end-to-end, hướng MediaPipe + CNN của đề tài có ưu điểm là phù hợp deadline hơn: MediaPipe giải quyết định vị mặt/mắt/miệng bằng mô hình pretrained, còn CNN tự xây tập trung vào phân loại trạng thái mắt. Điều này làm pipeline dễ giải thích, dễ deploy TFLite và phù hợp Android phone.",
            "Điểm khác biệt của bài làm là đặt trọng tâm vào prototype Android native: CameraX lấy frame realtime, MediaPipe trích landmark, EAR/MAR làm baseline ổn định, CNN/TFLite bổ sung thành phần học sâu và event log chỉ ghi trạng thái cảnh báo thay vì lưu video khuôn mặt.",
        ],
    )


def add_chapter_2(doc: Document) -> None:
    doc.add_heading("Chương 2: Chuẩn bị dữ liệu", level=1)
    doc.add_heading("2.1. Thu thập dữ liệu", level=2)
    add_paragraphs(
        doc,
        [
            "Dữ liệu chính của đề tài là bộ MRL Eye Dataset và các bản fork/tiền xử lý của MRL đã có trong thư mục rawdata. Dữ liệu gồm ảnh mắt hồng ngoại được chia thành trạng thái mắt mở và mắt nhắm, phù hợp với bài toán phân loại eye-state.",
            "Các thư mục raw hiện có gồm mrleyedataset/Open-Eyes, mrleyedataset/Close-Eyes, rawdata/open_eye, rawdata/closed_eye và rawdata/data đã chia sẵn train/val/test với nhãn awake/sleepy. Trong quá trình chuẩn hóa, awake được ánh xạ thành eyes_open và sleepy được ánh xạ thành eyes_closed.",
            "Ngoài dữ liệu public, nhóm có thể tự quay video bằng điện thoại Android để kiểm thử thực tế. Dữ liệu tự quay nên ưu tiên dùng cho đánh giá demo, không trộn hết vào train để tránh kết quả đánh giá bị lạc quan quá mức.",
        ],
    )
    add_dataset_inventory_table(doc)
    add_paragraphs(
        doc,
        [
            "Dữ liệu được tổ chức về cấu trúc chuẩn dataset/train, dataset/val, dataset/test, mỗi split gồm hai lớp eyes_open và eyes_closed. Script tools/prepare_eye_dataset.py được dùng để ánh xạ rawdata/data sang cấu trúc chuẩn, sau đó tools/summarize_dataset.py sinh bảng thống kê và manifest cho báo cáo.",
        ],
    )

    doc.add_heading("2.2. Tiền xử lý dữ liệu", level=2)
    add_paragraphs(
        doc,
        [
            "Các bước tiền xử lý chính gồm chuẩn hóa nhãn, chia dữ liệu theo train/validation/test, resize ảnh về kích thước 64x64, chuyển ảnh về RGB nếu cần và chuẩn hóa giá trị pixel về khoảng [0,1]. Đây là các bước cơ bản đã học/thực hành trong xử lý dữ liệu ảnh.",
            "Kỹ thuật chưa học sâu trên lớp là sử dụng MediaPipe Face Landmarker để trích landmarks và crop vùng mắt làm ROI trước khi đưa vào CNN. Cách làm này giúp mô hình CNN tập trung vào vùng thông tin quan trọng, giảm nhiễu từ nền, tóc, mũ hoặc các vùng không liên quan.",
            "Việc tiền xử lý là cần thiết vì dữ liệu ảnh có thể khác nhau về kích thước, độ sáng, định dạng và cách đặt nhãn. Nếu không chuẩn hóa, mô hình có thể học sai đặc trưng hoặc không tương thích với input của TensorFlow Lite trên Android.",
        ],
    )
    add_placeholder(doc, "Chèn hình minh họa trước/sau tiền xử lý: ảnh gốc, vùng mắt crop, ảnh resize 64x64.")


def add_chapter_3(doc: Document) -> None:
    doc.add_heading("Chương 3: Xây dựng mô hình", level=1)
    doc.add_heading("3.1. Trích chọn đặc trưng", level=2)
    add_paragraphs(
        doc,
        [
            "Dữ liệu của đề tài là ảnh/video nên đặc trưng được trích theo hai hướng. Hướng thứ nhất là đặc trưng hình học từ landmarks khuôn mặt: EAR cho mắt và MAR cho miệng. Hướng thứ hai là đặc trưng học tự động bằng CNN từ vùng ảnh mắt.",
            "EAR được tính từ các điểm landmark quanh mắt để phản ánh độ mở của mắt. Khi mắt nhắm, khoảng cách theo chiều dọc giảm khiến EAR giảm. MAR được tính từ landmarks quanh miệng để phát hiện há miệng/ngáp. Các đặc trưng này dễ giải thích và phù hợp làm baseline realtime.",
            "CNN không cần thiết kế đặc trưng thủ công mà học bộ lọc tích chập từ ảnh mắt. Trong pipeline, MediaPipe giải quyết bài toán định vị mắt, còn CNN giải quyết bài toán phân loại eyes_open/eyes_closed.",
        ],
    )

    doc.add_heading("3.2. Lựa chọn thuật toán/mô hình", level=2)
    add_paragraphs(
        doc,
        [
            "Đề tài chọn mô hình chính là MediaPipe Face Landmarker kết hợp CNN nhỏ. MediaPipe là mô hình pretrained được dùng để phát hiện landmarks khuôn mặt. CNN là mô hình tự xây bằng TensorFlow/Keras để phân loại ảnh mắt thành hai lớp eyes_open và eyes_closed.",
            "Dữ liệu được chia theo train/validation/test. Vì số lượng ảnh hai lớp tương đối cân bằng, đề tài chưa cần áp dụng kỹ thuật cân bằng dữ liệu phức tạp. Nếu về sau kết hợp thêm dataset khác gây lệch lớp, có thể dùng class weighting hoặc lấy mẫu cân bằng.",
            "YOLO end-to-end được khảo sát như hướng mở rộng. Tuy nhiên YOLO cần bounding box labels cho mắt mở, mắt nhắm hoặc ngáp; đồng thời quá trình export và hậu xử lý trên Android phức tạp hơn. Vì vậy, với phạm vi bài tập lớn, MediaPipe + CNN là lựa chọn phù hợp hơn.",
        ],
    )

    doc.add_heading("3.3. Cấu hình huấn luyện mô hình", level=2)
    add_training_config_table(doc)
    add_paragraphs(
        doc,
        [
            "Mô hình CNN gồm các lớp Conv2D, MaxPooling2D, GlobalAveragePooling2D, Dropout và Dense Softmax. Mô hình được train bằng TensorFlow/Keras, sau đó export sang TensorFlow Lite để Android app load trực tiếp trong thư mục app/src/main/assets.",
            "Hyper-parameter mặc định trong script gồm image size 64x64, batch size 32, số epoch 12, optimizer Adam với learning rate 1e-3. Các tham số này đủ nhẹ để train nhanh và phù hợp thiết bị Android khi inference.",
        ],
    )


def add_chapter_4(doc: Document) -> None:
    doc.add_heading("Chương 4: Đánh giá mô hình và demo", level=1)
    doc.add_heading("4.1. Độ đo đánh giá", level=2)
    add_paragraphs(
        doc,
        [
            "Các độ đo đánh giá gồm accuracy, precision, recall, F1-score và confusion matrix. Accuracy cho biết tỷ lệ dự đoán đúng tổng thể. Precision cho biết khi mô hình dự đoán một lớp thì mức độ đúng là bao nhiêu. Recall cho biết mô hình phát hiện được bao nhiêu mẫu thật của lớp đó.",
            "Với bài toán cảnh báo buồn ngủ, recall của lớp eyes_closed đặc biệt quan trọng vì bỏ sót trạng thái nhắm mắt kéo dài có thể nguy hiểm hơn báo nhầm nhẹ. F1-score được dùng để cân bằng precision và recall.",
        ],
    )
    add_metrics_formula_table(doc)

    doc.add_heading("4.2. Kết quả đánh giá", level=2)
    add_note(
        doc,
        "Mục này cần điền số liệu thật sau khi chạy tools/evaluate_eye_classifier.py. "
        "Các placeholder dưới đây được giữ để bạn thay bằng kết quả cuối.",
    )
    add_results_table(doc)
    add_placeholder(doc, "Chèn confusion matrix từ outputs/evaluation/confusion_matrix.csv.")
    add_placeholder(doc, "Chèn biểu đồ hoặc ảnh minh họa một số dự đoán đúng/sai.")

    doc.add_heading("4.3. Demo", level=2)
    add_paragraphs(
        doc,
        [
            "Prototype chính là Android native app. App sử dụng CameraX để lấy frame realtime, MediaPipe Face Landmarker để trích landmarks, DrowsinessAnalyzer để tính EAR/MAR và TfliteDrowsinessClassifier để chạy model CNN nếu drowsiness_model.tflite tồn tại.",
            "Giao diện demo hiển thị camera preview, overlay gồm EAR, MAR, FPS, trạng thái Awake/Drowsy/Yawning và nhãn CNN nếu model load được. Khi mắt nhắm kéo dài, app phát cảnh báo bằng âm thanh/rung và ghi event log tối thiểu.",
            "Ngoài Android, project có streamlit_app.py làm fallback demo. Streamlit phù hợp để trình bày ảnh upload/camera snapshot, pipeline AI và kết quả đánh giá; tuy nhiên Android native vẫn là hướng deploy chính vì chạy trực tiếp trên thiết bị tài xế.",
        ],
    )
    add_placeholder(doc, "Chèn ảnh giao diện Android app khi trạng thái Awake.")
    add_placeholder(doc, "Chèn ảnh/video demo khi trạng thái Drowsy alert hoặc Yawning.")
    add_placeholder(doc, "Chèn ảnh Streamlit fallback nếu dùng trong phần demo.")


def add_conclusion(doc: Document) -> None:
    doc.add_heading("Kết luận", level=1)
    add_paragraphs(
        doc,
        [
            "Đề tài đã xây dựng được pipeline phát hiện buồn ngủ tài xế theo hướng MediaPipe + EAR/MAR + CNN/TFLite. Hệ thống có khả năng lấy frame từ camera, xác định landmarks khuôn mặt, phân tích trạng thái mắt/miệng, cảnh báo theo thời gian và định hướng deploy trên Android phone.",
            "Những điểm đã làm được gồm tổ chức project Android native, chuẩn bị pipeline train/evaluate CNN, chuẩn hóa dataset eye-state, tạo tài liệu triển khai Android/Streamlit và xây dựng khung báo cáo bám rubric IS54A.",
            "Hạn chế hiện tại gồm phụ thuộc ánh sáng, góc camera, kính/che mặt; dữ liệu tự quay còn nhỏ; chưa được kiểm thử trên nhiều dòng camera hành trình; và chưa có chứng nhận an toàn để dùng như sản phẩm thực tế. Hướng phát triển là bổ sung PERCLOS, head pose, gaze estimation, YOLO end-to-end, tối ưu TFLite/ONNX và triển khai trên dashcam thông minh hoặc thiết bị embedded trong xe.",
        ],
    )


def add_references(doc: Document) -> None:
    doc.add_heading("Tài liệu tham khảo", level=1)
    refs = [
        "MediaPipe Face Landmarker Android. https://ai.google.dev/edge/mediapipe/solutions/vision/face_landmarker/android",
        "Android CameraX ImageAnalysis. https://developer.android.com/training/camerax/analyze",
        "TensorFlow Lite Inference Guide. https://www.tensorflow.org/lite/guide/inference",
        "Streamlit Community Cloud Deployment. https://docs.streamlit.io/deploy/streamlit-community-cloud/deploy-your-app/deploy",
        "Ultralytics YOLO Export Documentation. https://docs.ultralytics.com/modes/export/",
        "MRL Eye Dataset. http://mrl.cs.vsb.cz/eyedataset",
        "NTHU Driver Drowsiness Detection Dataset. https://cv.cs.nthu.edu.tw/php/callforpaper/datasets/DDD/",
        "YawDD: Yawning Detection Dataset. https://ieee-dataport.org/open-access/yawdd-yawning-detection-dataset",
    ]
    for ref in refs:
        p = doc.add_paragraph(style=None)
        p.style = doc.styles["Normal"]
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.add_run(ref)


def add_dataset_inventory_table(doc: Document) -> None:
    doc.add_paragraph("Bảng 2.1. Kiểm kê dữ liệu raw hiện có trong project.")
    rows = [
        ("mrleyedataset/Close-Eyes", "Close-Eyes", "eyes_closed", "41,946", "MRL raw class folder"),
        ("mrleyedataset/Open-Eyes", "Open-Eyes", "eyes_open", "42,952", "MRL raw class folder"),
        ("rawdata/closed_eye", "closed_eye", "eyes_closed", "24,000", "Flattened class folder, có khả năng là subset/copy"),
        ("rawdata/open_eye", "open_eye", "eyes_open", "24,000", "Flattened class folder, có khả năng là subset/copy"),
        ("rawdata/data/train/awake", "awake", "eyes_open", "25,770", "Pre-split MRL train"),
        ("rawdata/data/train/sleepy", "sleepy", "eyes_closed", "25,167", "Pre-split MRL train"),
        ("rawdata/data/val/awake", "awake", "eyes_open", "8,591", "Pre-split MRL validation"),
        ("rawdata/data/val/sleepy", "sleepy", "eyes_closed", "8,389", "Pre-split MRL validation"),
        ("rawdata/data/test/awake", "awake", "eyes_open", "8,591", "Pre-split MRL test"),
        ("rawdata/data/test/sleepy", "sleepy", "eyes_closed", "8,390", "Pre-split MRL test"),
    ]
    add_table(doc, ["Nguồn", "Nhãn gốc", "Nhãn chuẩn", "Số ảnh", "Ghi chú"], rows, [2600, 1600, 1600, 1200, 3160])


def add_training_config_table(doc: Document) -> None:
    rows = [
        ("Ngôn ngữ train", "Python"),
        ("Framework train", "TensorFlow/Keras"),
        ("Deploy mobile", "TensorFlow Lite"),
        ("Input model", "Ảnh mắt 64x64x3"),
        ("Output model", "eyes_closed, eyes_open"),
        ("Batch size", "32"),
        ("Epoch", "12 hoặc điều chỉnh theo validation loss"),
        ("Optimizer", "Adam, learning rate 1e-3"),
        ("Phần cứng", "[ĐIỀN CPU/GPU/RAM máy train thực tế]"),
    ]
    add_table(doc, ["Thành phần", "Cấu hình"], rows, [2600, 6760])


def add_metrics_formula_table(doc: Document) -> None:
    rows = [
        ("Accuracy", "(TP + TN) / (TP + TN + FP + FN)", "Tỷ lệ dự đoán đúng tổng thể."),
        ("Precision", "TP / (TP + FP)", "Khi mô hình dự đoán một lớp, tỷ lệ đúng của dự đoán đó."),
        ("Recall", "TP / (TP + FN)", "Trong các mẫu thật của một lớp, mô hình phát hiện được bao nhiêu."),
        ("F1-score", "2 x Precision x Recall / (Precision + Recall)", "Trung bình điều hòa giữa precision và recall."),
        ("Confusion matrix", "Bảng đếm nhãn thật và nhãn dự đoán", "Cho thấy mô hình nhầm lớp nào với lớp nào."),
    ]
    add_table(doc, ["Độ đo", "Công thức/biểu diễn", "Ý nghĩa"], rows, [1800, 3000, 4560])


def add_results_table(doc: Document) -> None:
    rows = [
        ("CNN eye-state v1", "0.9721", "0.9608", "0.9838", "0.9721", "Metrics cho lớp eyes_closed; model tự xây, export TFLite."),
        ("EAR/MAR baseline", "[ĐIỀN SAU DEMO]", "[ĐIỀN SAU DEMO]", "[ĐIỀN SAU DEMO]", "[ĐIỀN SAU DEMO]", "Baseline hình học từ landmarks; đo sau khi chạy Android demo."),
    ]
    add_table(doc, ["Mô hình", "Accuracy", "Precision", "Recall", "F1-score", "Ghi chú"], rows, [2200, 1300, 1300, 1300, 1300, 2360])


def add_table(doc: Document, headers: list[str], rows: list[tuple[str, ...]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
        shade_cell(table.rows[0].cells[i], "D9EAF7")
        set_cell_text_bold(table.rows[0].cells[i])
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


def add_paragraphs(doc: Document, texts: list[str]) -> None:
    for text in texts:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(text)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.2)
    r = p.add_run("Ghi chú: ")
    r.bold = True
    p.add_run(text)
    shade_paragraph(p, "FFF3D8")


def add_placeholder(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run("[PLACEHOLDER] ")
    r.bold = True
    r.font.color.rgb = RGBColor(155, 28, 28)
    p.add_run(text).italic = True


def add_key_value(doc: Document, key: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.2)
    r = p.add_run(f"{key}: ")
    r.bold = True
    p.add_run(value)


def add_spacer(doc: Document, lines: int = 1) -> None:
    for _ in range(lines):
        doc.add_paragraph()


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Bấm chuột phải và chọn Update Field để cập nhật mục lục."
    fld_sep_run = OxmlElement("w:r")
    fld_sep_run.append(text)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    paragraph._p.append(fld_sep_run)
    run._r.append(fld_end)


def set_table_widths(table, widths: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.tcW
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[i]))
            tc_w.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_cell_text_bold(cell) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True


if __name__ == "__main__":
    main()
