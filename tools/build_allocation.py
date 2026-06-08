# -*- coding: utf-8 -*-
"""Sinh BẢNG PHÂN CÔNG CÔNG VIỆC (.docx) đúng format template — phủ mọi section báo cáo."""
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(13)

h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h.add_run("BẢNG PHÂN CÔNG CÔNG VIỆC"); r.bold = True; r.font.size = Pt(14)
doc.add_paragraph("Tổng phần trăm đóng góp của các thành viên = 100%. Mọi mục trong quyển báo cáo "
                  "(kể cả Tài liệu tham khảo) đều xuất hiện trong bảng dưới đây.").italic = True

COLS = ["Họ tên – Mã sinh viên", "Quyển báo cáo", "Sản phẩm (mã nguồn)", "% – Ký tên"]
WID = [Cm(3.2), Cm(6.2), Cm(6.8), Cm(2.3)]

members = [
    ("Đỗ Vũ Việt Hằng\nMSSV: 26A404____",
     ["Chương 2: CHUẨN BỊ DỮ LIỆU",
      "2.1. CNN phân loại trạng thái mắt (MRL)",
      "2.2. CNN phân loại trạng thái ngáp (Yawn)",
      "2.3. YOLOv11 (Datio_yolo)",
      "2.4. YOLO26", "2.5. RT-DETR", "2.6. RF-DETR",
      "Tài liệu tham khảo (Chương 2)"],
     ["Thu thập & tiền xử lý dữ liệu:",
      "- tools/pipeline_visual.py (10 bước tiền xử lý)",
      "- data_pipeline/preprocess.py",
      "- colab_extract_frames.ipynb (cắt frame video)",
      "- tools/step8_realdata_train.py (crop MediaPipe + trộn dataset)",
      "- tools/process_real_video.py",
      "- tools/build_report_ch2_full.js"],
     "25%"),

    ("(Họ tên) Quang\nMSSV: 26A404____",
     ["Chương 3: XÂY DỰNG MÔ HÌNH",
      "3.2. Lựa chọn thuật toán và mô hình",
      "3.3. Cấu hình huấn luyện mô hình",
      "3.4. So sánh hai phương pháp chia tập dữ liệu",
      "3.5. Xây dựng các mô hình phát hiện vật thể",
      "Tài liệu tham khảo (Chương 3)"],
     ["Xây dựng & huấn luyện mô hình:",
      "- train_cnn_eye.py / .ipynb (CNN mắt)",
      "- train_cnn_yawn.py / .ipynb (CNN ngáp)",
      "- colab_yolo11s.py, colab_yolo26_training.py",
      "- colab_rtdetr.py, colab_rfdetr_6class.ipynb",
      "- notebooks/AllInOne_Pipeline.ipynb",
      "- app/.../TfliteDrowsinessClassifier.kt, YawnClassifier.kt, YoloDetector.kt",
      "- tools/build_report_ch3.js"],
     "25%"),

    ("(Họ tên) Thu Hoài\nMSSV: 26A404____",
     ["Chương 1: GIỚI THIỆU (TỔNG QUAN)",
      "1.1. Đặt vấn đề và tính cấp thiết",
      "1.2. Tổng quan nghiên cứu liên quan và khoảng trống nghiên cứu",
      "1.3. Cơ chế cảnh báo và an toàn (âm thanh, thời gian lái)",
      "1.4. Mục tiêu, phạm vi và cấu trúc báo cáo",
      "Chương 3: 3.1. Trích chọn đặc trưng",
      "Tài liệu tham khảo (Chương 1)"],
     ["Trích chọn đặc trưng & cơ chế cảnh báo:",
      "- app/.../DrowsinessAnalyzer.kt (EAR/MAR từ landmark)",
      "- app/.../AlertController.kt (âm thanh binaural beta §1.3)",
      "- streamlit_app.py — _ear / _mar / _make_landmarker, make_binaural"],
     "25%"),

    ("Nguyễn Tuấn Đạt\nMSSV: 26A404____",
     ["Chương 4: ĐÁNH GIÁ VÀ SO SÁNH CÁC MÔ HÌNH",
      "4.1. Độ đo đánh giá",
      "4.2. Thực nghiệm và kết quả đánh giá",
      "4.3. Demo sản phẩm trên thiết bị Android",
      "Tài liệu tham khảo (Chương 4)"],
     ["Đánh giá, ứng dụng & triển khai:",
      "- tools/evaluate_combined.py (chỉ số đầy đủ)",
      "- evaluation_report.ipynb, colab_compare_all.ipynb",
      "- app/.../MainActivity.kt (tích hợp, fusion CNN-primary)",
      "- app/.../AnalyticsActivity.kt (báo cáo + khung giờ buồn ngủ)",
      "- app/.../EmergencyContactActivity.kt, EmergencyDispatcher.kt, UiKit.kt",
      "- streamlit_app.py (demo web), Dockerfile, STREAMLIT_DEPLOY.md",
      "- tools/build_report_ch4.js, tools/merge_report.py"],
     "25%"),
]

tbl = doc.add_table(rows=1, cols=4); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, c in enumerate(tbl.rows[0].cells):
    c.width = WID[i]
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COLS[i]); run.bold = True

def fill(cell, lines, center=False, bold_first=False):
    cell.paragraphs[0].text = ""
    for j, ln in enumerate(lines):
        p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
        if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(ln); run.font.size = Pt(11.5)
        if bold_first and j == 0: run.bold = True

for name, report, code, pct in members:
    row = tbl.add_row().cells
    for i in range(4):
        row[i].width = WID[i]
    fill(row[0], name.split("\n"), bold_first=True)
    fill(row[1], report, bold_first=True)
    fill(row[2], code, bold_first=True)
    fill(row[3], [pct, "", "(Ký tên)"], center=True)

doc.add_paragraph()
note = doc.add_paragraph()
note.add_run("Ghi chú: ").bold = True
note.add_run("Điền đầy đủ Họ tên + Mã sinh viên của Quang và Thu Hoài; tỉ lệ % điều chỉnh theo "
             "thực tế đóng góp sao cho tổng = 100%. Cấu trúc Chương 1 ghi theo đúng đề mục trong "
             "quyển báo cáo của nhóm.").italic = True

doc.save("outputs/PhanCong_CongViec.docx")
print("OK -> outputs/PhanCong_CongViec.docx")
