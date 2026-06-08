# -*- coding: utf-8 -*-
"""
HỢP NHẤT báo cáo thành 1 file Word:
  Ch2 (BaoCao_Chuong2.docx của tôi) + Ch3 + Ch4 (HoanChinh, bản tốt).
- Thêm mục "Demo sản phẩm trên thiết bị Android" + 3 ảnh demo vào Ch4 (trước phần Nhận xét).
- Thêm thụt đầu dòng 1.27cm cho thân bài Ch3 (HoanChinh Ch3 thiếu).
- Bỏ hẳn bản Ch3/Ch4 tôi sinh trước đó (gây trùng lặp + mâu thuẫn cấu hình).
Chạy: py -3.12 tools/merge_report.py
"""
import docx
from docx.shared import Cm, Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docxcompose.composer import Composer

DL = "C:/Users/ADMIN/Downloads"
CH2 = "outputs/BaoCao_Chuong2.docx"
CH3 = f"{DL}/Chuong3_DrowsyDriver_HoanChinh (1).docx"
CH4 = f"{DL}/Chuong4_DanhGiaMoHinh_HoanChinh.docx"
SHOTS = "outputs/demo_shots"
OUT = "outputs/BaoCao_HopNhat_DrowsyDriver.docx"

def insert_before(target, text="", style=None, align=None, italic=False, size=None):
    p = OxmlElement("w:p"); target._p.addprevious(p)
    para = Paragraph(p, target._parent)
    if style:
        try: para.style = style
        except Exception: pass
    if align is not None: para.alignment = align
    if text:
        r = para.add_run(text)
        if italic: r.italic = True
        if size: r.font.size = Pt(size)
    return para

def insert_image_before(target, path, width_in):
    p = OxmlElement("w:p"); target._p.addprevious(p)
    para = Paragraph(p, target._parent); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(path, width=Inches(width_in))
    return para

# ─── 1) Ch4: chèn mục Demo Android trước "Nhận xét tổng hợp" ───
d4 = docx.Document(CH4)
target = None
for p in d4.paragraphs:
    if p.style.name.startswith("Heading") and "Nhận xét" in p.text:
        target = p; break
if target is None:           # fallback: chèn cuối tài liệu
    target = d4.add_paragraph("")

demo = [
    ("h", "Demo sản phẩm trên thiết bị Android"),
    ("p", "Để minh chứng tính khả thi triển khai thực tế, hệ thống đã được đóng gói thành ứng dụng Android và chạy thử nghiệm trực tiếp trên điện thoại Sony Xperia (model SO-01L). Toàn bộ quá trình suy luận diễn ra trên thiết bị (on-device), không cần kết nối mạng. Các hình dưới đây minh họa giao diện và hoạt động thực tế của sản phẩm."),
    ("img", f"{SHOTS}/awake.png"),
    ("cap", "Hình 4.9. Màn hình giám sát chính: phát hiện trạng thái tỉnh táo theo thời gian thực (CNN làm chủ đạo)"),
    ("p", "Màn hình chính hiển thị luồng camera trước, trạng thái tài xế, các chỉ số EAR, MAR, độ tin cậy của mạng CNN và nguồn quyết định cảnh báo. Khi tài xế tỉnh táo, hệ thống hiển thị trạng thái màu xanh với độ tin cậy cao, đúng theo cơ chế hợp nhất lấy CNN làm chủ đạo và EAR/MAR làm lớp dự phòng."),
    ("img", f"{SHOTS}/analytics.png"),
    ("cap", "Hình 4.10. Màn hình báo cáo thống kê phiên lái: mức độ an toàn, chỉ số PERCLOS và số lần buồn ngủ"),
    ("p", "Màn hình báo cáo tổng hợp mức độ an toàn của phiên lái, thời gian lái, số lần buồn ngủ và ngáp, chỉ số PERCLOS (tỉ lệ thời gian nhắm mắt) cùng trạng thái của tính năng tự động gửi định vị, giúp tài xế tự theo dõi mức độ mệt mỏi của bản thân."),
    ("img", f"{SHOTS}/emergency.png"),
    ("cap", "Hình 4.11. Màn hình liên hệ khẩn cấp: tự động gửi vị trí cho người thân khi vượt ngưỡng buồn ngủ"),
    ("p", "Màn hình liên hệ khẩn cấp cho phép lưu thông tin người thân và đặt ngưỡng số lần buồn ngủ cho mỗi chuyến đi. Khi tài xế vượt ngưỡng, ứng dụng tự động gửi tin nhắn SMS kèm vị trí GPS (liên kết Google Maps) cho người thân mà không cần thao tác, tăng khả năng ứng cứu kịp thời."),
]
import os
for kind, val in demo:
    if kind == "h":
        insert_before(target, val, style="Heading 2")
    elif kind == "p":
        pp = insert_before(target, val, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        pp.paragraph_format.first_line_indent = Cm(1.27)
    elif kind == "img":
        if os.path.exists(val): insert_image_before(target, val, 2.3)
    elif kind == "cap":
        insert_before(target, val, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=12)
d4.save("outputs/_merge_ch4.docx")
print("Ch4 + demo Android: OK")

# ─── 2) Ch3: thêm thụt đầu dòng 1.27cm cho thân bài ───
d3 = docx.Document(CH3)
cnt = 0
for p in d3.paragraphs:
    if p.style.name != "Normal": continue
    t = p.text.strip()
    if not t: continue
    if p.alignment == WD_ALIGN_PARAGRAPH.CENTER: continue
    if t[:5] in ("Hình ", "Bảng ") or t.startswith("Nguồn"): continue
    p.paragraph_format.first_line_indent = Cm(1.27)
    cnt += 1
d3.save("outputs/_merge_ch3.docx")
print(f"Ch3 + thụt lề 1.27cm: {cnt} đoạn")

# ─── 3) Merge Ch2 + Ch3 + Ch4 ───
master = docx.Document(CH2)
comp = Composer(master)
master.add_page_break()
comp.append(docx.Document("outputs/_merge_ch3.docx"))
master.add_page_break()
comp.append(docx.Document("outputs/_merge_ch4.docx"))

# ─── 3b) ĐÁNH SỐ CỨNG vào text heading (auto-number xung đột sau merge) ───
import string
def sname(p):
    try: return p.style.name or ""
    except Exception: return ""

LVL = {"Heading 1": 0, "Heading 2": 1, "Heading 3": 2, "Heading 4": 3}
ch = 1; sec = 0; sub = 0; let = 0; fixed = 0
for p in master.paragraphs:
    sn = sname(p)
    if sn not in LVL or not p.text.strip():
        continue
    # bỏ auto-numbering cũ (numPr) để không nhân đôi số
    pPr = p._p.pPr
    if pPr is not None:
        npr = pPr.find(qn("w:numPr"))
        if npr is not None: pPr.remove(npr)
    lvl = LVL[sn]
    if lvl == 0:
        ch += 1; sec = sub = let = 0; label = f"CHƯƠNG {ch}: "
    elif lvl == 1:
        sec += 1; sub = let = 0; label = f"{ch}.{sec}. "
    elif lvl == 2:
        sub += 1; let = 0; label = f"{ch}.{sec}.{sub}. "
    else:
        let += 1; label = f"{string.ascii_lowercase[(let-1) % 26]}. "
    if p.runs:
        p.runs[0].text = label + p.runs[0].text.lstrip()
    else:
        p.add_run(label)
    fixed += 1
print(f"Đánh số cứng heading: {fixed} heading (Ch2..Ch{ch})")

comp.save(OUT)
print("ĐÃ HỢP NHẤT ->", OUT)
